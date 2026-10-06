#!/usr/bin/env python3
"""Build a PowerPoint deck with one full-slide image per PNG."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

EXAMPLE = """Examples:
  pngs_to_pptx.py -o pic/slides.pptx pic/a.png pic/b.png
  pngs_to_pptx.py -o pic/deck.pptx --manifest pic/ppt_manifest.tsv pic/a.png pic/b.png
  pngs_to_pptx.py -o pic/deck.pptx --run-dir /path/to/run --invocation 'airport_comment_plots.bash ...' pic/a.png
"""


def create_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Write a .pptx with one slide per PNG (image scaled to fit, optional footer text).",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=EXAMPLE,
    )
    parser.add_argument(
        "images",
        nargs="+",
        metavar="PNG",
        help="PNG files in slide order",
    )
    parser.add_argument(
        "-o",
        "--output",
        required=True,
        metavar="FILE",
        help="Output PowerPoint path (.pptx)",
    )
    parser.add_argument(
        "--manifest",
        metavar="FILE",
        default=None,
        help="TSV manifest: PNG_PATH<TAB>footer command (one plot per line).",
    )
    parser.add_argument(
        "--footer-font-size",
        type=int,
        default=8,
        metavar="PT",
        help="Footer command font size in points (default: %(default)s).",
    )
    parser.add_argument(
        "--run-dir",
        default=None,
        metavar="DIR",
        help="Working directory for final info slide.",
    )
    parser.add_argument(
        "--invocation",
        default=None,
        help="Full airport_comment_plots.bash command for final info slide.",
    )
    return parser


def load_manifest(path: Path) -> dict[str, str]:
    captions: dict[str, str] = {}
    text = path.read_text(encoding="utf-8")
    for line in text.splitlines():
        if not line.strip():
            continue
        png, sep, cap = line.partition("\t")
        if not sep:
            continue
        key = str(Path(png).resolve())
        captions[key] = cap
        captions[str(Path(png))] = cap
    return captions


def add_footer(slide, text: str, slide_w: int, slide_h: int, font_pt: int) -> None:
    if not text:
        return
    from pptx.util import Inches, Pt

    margin = Inches(0.08)
    box_h = Inches(0.45)
    left = margin
    top = slide_h - box_h - margin
    width = slide_w - 2 * margin
    tx = slide.shapes.add_textbox(left, top, width, box_h)
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_pt)


def fit_picture_on_slide(
    slide,
    png_path: Path,
    slide_w: int,
    slide_h: int,
    footer_reserve: int,
) -> None:
    from PIL import Image

    with Image.open(png_path) as im:
        iw, ih = im.size
    if iw <= 0 or ih <= 0:
        raise ValueError(f"invalid image size: {png_path}")
    avail_h = slide_h - footer_reserve
    scale = min(slide_w / iw, avail_h / ih)
    w = int(iw * scale)
    h = int(ih * scale)
    left = (slide_w - w) // 2
    top = max(0, (avail_h - h) // 2)
    slide.shapes.add_picture(str(png_path), left, top, width=w, height=h)


def add_info_slide(prs, run_dir: str | None, invocation: str | None, font_pt: int) -> None:
    from pptx.util import Inches, Pt

    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    sw, sh = prs.slide_width, prs.slide_height
    parts = []
    if run_dir:
        parts.append(f"Directory: {run_dir}")
    if invocation:
        parts.append(invocation.strip())
    body = "\n\n".join(parts)
    if not body:
        return
    margin = Inches(0.5)
    tx = slide.shapes.add_textbox(margin, margin, sw - 2 * margin, sh - 2 * margin)
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = body
    p.font.size = Pt(font_pt)


def build_pptx(
    png_paths: list[Path],
    out_path: Path,
    captions: dict[str, str],
    footer_font_pt: int,
    run_dir: str | None,
    invocation: str | None,
) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise SystemExit(
            "python-pptx is required: pip install python-pptx"
        ) from exc

    footer_reserve = Inches(0.5)
    prs = Presentation()
    blank = prs.slide_layouts[6]
    sw, sh = prs.slide_width, prs.slide_height
    for png in png_paths:
        slide = prs.slides.add_slide(blank)
        fit_picture_on_slide(slide, png, sw, sh, footer_reserve)
        cap = captions.get(str(png.resolve()), captions.get(str(png), ""))
        add_footer(slide, cap, sw, sh, footer_font_pt)
    if run_dir or invocation:
        add_info_slide(prs, run_dir, invocation, footer_font_pt)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))


def main(argv: list[str] | None = None) -> int:
    parser = create_parser()
    args = parser.parse_args(argv)
    out_path = Path(args.output).resolve()
    captions: dict[str, str] = {}
    if args.manifest:
        man = Path(args.manifest)
        if not man.is_file():
            print(f"Error: manifest not found: {man}", file=sys.stderr)
            return 1
        captions = load_manifest(man)

    png_paths: list[Path] = []
    for raw in args.images:
        p = Path(raw)
        if not p.is_file():
            print(f"Error: PNG not found: {p}", file=sys.stderr)
            return 1
        if p.suffix.lower() != ".png":
            print(f"Error: not a PNG file: {p}", file=sys.stderr)
            return 1
        png_paths.append(p.resolve())
    if not png_paths:
        print("Error: no PNG files given", file=sys.stderr)
        return 1

    build_pptx(
        png_paths,
        out_path,
        captions,
        args.footer_font_size,
        args.run_dir,
        args.invocation,
    )
    n_slides = len(png_paths) + (1 if (args.run_dir or args.invocation) else 0)
    print(f"saved PowerPoint: {out_path} ({n_slides} slides)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
